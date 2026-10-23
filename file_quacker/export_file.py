"""Export a DuckDB source to a flat file via DuckDB's ``COPY ... TO``.

Flat (any delimiter) + Parquet are single-statement atomic writes via
DuckDB; Excel uses ``openpyxl`` streaming (DuckDB's native xlsx support
requires the ``spatial`` extension, which we don't pull in).

Every exporter respects the dialog's column mapping: source columns are
projected through ``"source" as "target"`` aliases, so the user's target
renames appear in the file's header / schema.
"""

from __future__ import annotations

import contextlib
import time
from dataclasses import dataclass
from threading import RLock
from typing import Literal

from . import db
from .export_source import ExportSource
from .text_values import string_columns, stripped_sql


@dataclass
class Mapping:
    source: str
    target: str


@dataclass
class FlatTarget:
    path: str
    delimiter: str = '|'
    header: bool = True
    quote: str = '"'
    null: str = ''
    # Default to CRLF; the most common destination is Windows tooling
    # / SQL Server bcp / Excel, all of which prefer it.  Unix consumers
    # tolerate CRLF; Windows consumers often choke on bare LF.
    line_ending: Literal['crlf', 'lf'] = 'crlf'


@dataclass
class ParquetTarget:
    path: str
    compression: Literal['snappy', 'zstd', 'gzip', 'uncompressed'] = 'snappy'


@dataclass
class ExcelTarget:
    path: str
    sheet_name: str = 'Sheet1'


FileTarget = FlatTarget | ParquetTarget | ExcelTarget


def _project(source: str, is_string: bool, trim_strings: bool) -> str:
    """SELECT expression for one source column. String columns get wrapped
    in regexp_replace when trim is on; an empty result is then promoted
    to NULL via nullif so a 25-spaces cell lands as NULL rather than as
    an empty string – consistent with how a truly-empty cell already
    behaves on the source side. Non-string columns are the bare ident."""
    qident = db.quote_ident(source)
    if trim_strings and is_string:
        return stripped_sql(qident)
    return qident


_lock = RLock()
_progress: dict = {'phase': 'idle', 'rows_written': 0, 'elapsed_ms': 0, 'error': None}
_cancel_flag = False
_running = False


def _set(**kw):
    with _lock:
        _progress.update(kw)


def get_progress() -> dict:
    with _lock:
        return dict(_progress)


def cancel() -> bool:
    global _cancel_flag
    _cancel_flag = True
    return True


def export_to_file(source: ExportSource | dict, target: dict,
                   mappings: list[dict] | None = None,
                   trim_strings: bool = True) -> dict:
    """Dispatch to the right exporter based on ``target['kind']``.

    ``trim_strings`` strips leading / trailing whitespace from every VARCHAR
    source column on the way out (default on). Non-string columns are
    untouched.
    """
    global _cancel_flag, _running
    with _lock:
        if _running:
            return {'ok': False, 'error': 'a file export is already running', 'elapsed_ms': 0}
        _running = True
    _cancel_flag = False
    if isinstance(source, dict):
        source = ExportSource(**source)
    maps = _coerce_mappings(mappings)
    kind = target.get('kind')
    try:
        if kind == 'flat':
            return _export_flat(source, FlatTarget(**_strip_kind(target)), maps, trim_strings)
        if kind == 'parquet':
            return _export_parquet(source, ParquetTarget(**_strip_kind(target)), maps, trim_strings)
        if kind == 'excel':
            return _export_excel(source, ExcelTarget(**_strip_kind(target)), maps, trim_strings)
        raise ValueError(f'unknown target kind: {kind!r}')
    finally:
        with _lock:
            _running = False


def _coerce_mappings(mappings: list[dict] | None) -> list[Mapping]:
    if not mappings:
        return []
    out: list[Mapping] = []
    for m in mappings:
        if isinstance(m, Mapping):
            out.append(m)
        else:
            out.append(Mapping(source=m['source'], target=m.get('target') or m['source']))
    return out


def _strip_kind(d: dict) -> dict:
    return {k: v for k, v in d.items() if k != 'kind'}


def _naive_time_columns(qtable: str) -> dict[str, str]:
    """Time-zone-aware columns on a table or view, mapped to the type
    that drops the zone. The cast keeps the session's local wall time
    (TIMETZ keeps its own), which is what the CSV export writes minus
    the offset."""
    naive = {'TIMESTAMP WITH TIME ZONE': 'timestamp', 'TIME WITH TIME ZONE': 'time'}
    info = db.conn().execute(f"""
        pragma table_info({qtable})
        ;
    """).fetchall()
    return {row[1]: naive[row[2]] for row in info if row[2] in naive}


def _select_with_mappings(qtable: str, mappings: list[Mapping],
                          trim_strings: bool = False,
                          naive_times: bool = False) -> tuple[str, list[str]]:
    """Return ``(select_list, target_column_names)`` honoring the user's
    source→target renames.  Falls back to every non-meta column on the
    table when mappings is empty. When ``trim_strings`` is on, VARCHAR
    columns get wrapped in regexp_replace so leading / trailing whitespace
    is stripped on the way out. ``naive_times`` casts time-zone-aware
    columns to their plain types, for writers that reject zones.
    """
    str_cols = string_columns(qtable) if trim_strings else set()
    tz_cols = _naive_time_columns(qtable) if naive_times else {}

    def expr_for(source: str) -> str:
        if source in tz_cols:
            return f'cast({db.quote_ident(source)} as {tz_cols[source]})'
        return _project(source, source in str_cols, trim_strings)

    if mappings:
        parts = []
        for m in mappings:
            parts.append(f'{expr_for(m.source)} as {db.quote_ident(m.target)}')
        targets = [m.target for m in mappings]
        return ', '.join(parts), targets

    # No mappings provided; use every column except _src_row_num, with
    # source names as both source AND target.
    c = db.conn()
    info = c.execute(f"""
        pragma table_info({qtable})
        ;
    """).fetchall()
    cols = [r[1] for r in info if r[1] != '_src_row_num']
    parts = []
    for n in cols:
        expr = expr_for(n)
        # Trimmed or cast columns need an alias to keep the column name in
        # the output; bare-ident projections don't.
        parts.append(expr if expr == db.quote_ident(n) else f'{expr} as {db.quote_ident(n)}')
    return ', '.join(parts), cols


def _export_flat(source: ExportSource, tgt: FlatTarget,
                 mappings: list[Mapping], trim_strings: bool = False) -> dict:
    t0 = time.perf_counter()
    _set(phase='writing', rows_written=0, elapsed_ms=0, error=None)
    src_name, cleanup = source.materialize()
    try:
        qtable = db.quote_ident(src_name)
        select_list, _ = _select_with_mappings(qtable, mappings, trim_strings)
        # Parameter placeholders aren't supported inside COPY's option list,
        # so we literal-escape values.  Path is a trusted local path typed
        # by the user in their own app; single quotes are doubled defensively.
        path = tgt.path.replace("'", "''")
        delim = tgt.delimiter.replace("'", "''")
        quote = tgt.quote.replace("'", "''")
        null = tgt.null.replace("'", "''")
        header = 'true' if tgt.header else 'false'
        # DuckDB's `new_line` accepts '\r\n', '\n', or '\r' (the literal
        # escape sequences inside the SQL string).  CRLF is the default
        # since most downstream tooling (Excel, SQL Server bcp, Windows
        # text consumers) expects it; LF stays available for Unix-y
        # pipelines.
        new_line = '\\r\\n' if tgt.line_ending == 'crlf' else '\\n'
        copy_sql = f"""
            copy (select {select_list} from {qtable})
            to '{path}'
            (header {header}, delimiter '{delim}',
             quote '{quote}', null '{null}',
             new_line '{new_line}')
            ;
        """
        c = db.conn()
        c.execute(copy_sql)
        (n,) = c.execute(f"""
            select  count(*)
            from    {qtable}
            ;
        """).fetchone()
        elapsed = int((time.perf_counter() - t0) * 1000)
        _set(phase='done', rows_written=int(n), elapsed_ms=elapsed)
        return {'ok': True, 'rows_written': int(n), 'path': tgt.path,
                'elapsed_ms': elapsed}
    except Exception as ex:
        _set(phase='error', error=str(ex))
        return {'ok': False, 'error': str(ex),
                'elapsed_ms': int((time.perf_counter() - t0) * 1000)}
    finally:
        cleanup()


def _export_parquet(source: ExportSource, tgt: ParquetTarget,
                    mappings: list[Mapping], trim_strings: bool = False) -> dict:
    t0 = time.perf_counter()
    _set(phase='writing', rows_written=0, elapsed_ms=0, error=None)
    src_name, cleanup = source.materialize()
    try:
        qtable = db.quote_ident(src_name)
        select_list, _ = _select_with_mappings(qtable, mappings, trim_strings)
        path = tgt.path.replace("'", "''")
        compression = tgt.compression.upper()
        copy_sql = f"""
            copy (select {select_list} from {qtable})
            to '{path}'
            (format parquet, compression {compression})
            ;
        """
        c = db.conn()
        c.execute(copy_sql)
        (n,) = c.execute(f"""
            select  count(*)
            from    {qtable}
            ;
        """).fetchone()
        elapsed = int((time.perf_counter() - t0) * 1000)
        _set(phase='done', rows_written=int(n), elapsed_ms=elapsed)
        return {'ok': True, 'rows_written': int(n), 'path': tgt.path,
                'elapsed_ms': elapsed}
    except Exception as ex:
        _set(phase='error', error=str(ex))
        return {'ok': False, 'error': str(ex),
                'elapsed_ms': int((time.perf_counter() - t0) * 1000)}
    finally:
        cleanup()


def _export_excel(source: ExportSource, tgt: ExcelTarget,
                  mappings: list[Mapping], trim_strings: bool = False) -> dict:
    import openpyxl  # lazy; only loaded when Excel export is used

    t0 = time.perf_counter()
    _set(phase='writing', rows_written=0, elapsed_ms=0, error=None)
    src_name, cleanup = source.materialize()
    ws = None
    try:
        qtable = db.quote_ident(src_name)
        # openpyxl rejects time-zone-aware datetimes.
        select_list, target_cols = _select_with_mappings(qtable, mappings, trim_strings,
                                                         naive_times=True)
        c = db.conn()
        cur = c.execute(f"""
            select  {select_list}
            from    {qtable}
            ;
        """)

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title=tgt.sheet_name)
        ws.append(target_cols)

        n = 0
        for row in cur.fetchall():
            if _cancel_flag:
                _set(phase='error', error='cancelled')
                return {'ok': False, 'error': 'cancelled',
                        'rows_written': n,
                        'elapsed_ms': int((time.perf_counter() - t0) * 1000)}
            ws.append(list(row))
            n += 1
            if n % 5000 == 0:
                _set(rows_written=n,
                     elapsed_ms=int((time.perf_counter() - t0) * 1000))

        wb.save(tgt.path)
        elapsed = int((time.perf_counter() - t0) * 1000)
        _set(phase='done', rows_written=n, elapsed_ms=elapsed)
        return {'ok': True, 'rows_written': n, 'path': tgt.path,
                'elapsed_ms': elapsed}
    except Exception as ex:
        _set(phase='error', error=str(ex))
        return {'ok': False, 'error': str(ex),
                'elapsed_ms': int((time.perf_counter() - t0) * 1000)}
    finally:
        # An unsaved write-only sheet keeps its temp file open, and
        # openpyxl's exit hook then fails to delete it. Best-effort, so
        # it can't mask the error that got us here.
        if ws is not None and not ws.closed:
            with contextlib.suppress(Exception):
                ws.close()
        cleanup()
