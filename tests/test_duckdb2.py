"""DuckDB 2.0 behavior changes FQ smooths over: division by zero, untyped
NULL columns, TIMESTAMPTZ fetches, and Esc-to-cancel."""

from __future__ import annotations

import os
import threading
import time

import openpyxl
import pytest

from file_quacker import db, ddl, derive, ingest, profile, sql
from file_quacker.api import Api
from file_quacker.export_file import export_to_file
from file_quacker.export_source import ExportSource


def _types(name: str) -> dict[str, str]:
    rows = db.conn().execute(f"""
        pragma table_info({db.quote_ident(name)})
        ;
    """).fetchall()
    return {r[1]: r[2] for r in rows}


def _ddl_columns(name: str, dialect: str) -> list[list[str]]:
    """Column lines of the generated DDL, split on whitespace."""
    lines = ddl.generate(name, dialect).splitlines()
    return [line.lstrip('(, ').split() for line in lines[1:-2]]


def _target(kind: str, path) -> dict:
    return {
        'flat': {'kind': 'flat', 'path': str(path), 'delimiter': ',', 'header': True,
                 'quote': '"', 'null': ''},
        'parquet': {'kind': 'parquet', 'path': str(path), 'compression': 'snappy'},
        'excel': {'kind': 'excel', 'path': str(path), 'sheet_name': 'Data'},
    }[kind]


_EXT = {'flat': 'csv', 'parquet': 'parquet', 'excel': 'xlsx'}


# --------------------------------------------------------------------------- #
# Division by zero                                                            #
# --------------------------------------------------------------------------- #

def test_division_by_zero_is_null_on_every_cursor():
    assert db.conn().execute('select 1 // 0, 5 % 0;').fetchone() == (None, None)

    out: list = []
    t = threading.Thread(target=lambda: out.append(
        db.conn().execute('select 7 // 0, 7 % 0;').fetchone()))
    t.start()
    t.join(10)
    assert out == [(None, None)]


def test_division_by_zero_is_null_in_sql_editor():
    res = sql.run('select 1 // 0 as a, 5 % 0 as b;')
    assert res['ok'] is True, res
    assert res['rows'] == [[None, None]]


# --------------------------------------------------------------------------- #
# Untyped NULL columns                                                        #
# --------------------------------------------------------------------------- #

def test_sql_run_result_has_no_null_types():
    res = sql.run("""
        select  null            as a
        ,       [null]          as b
        ,       {'k': null}     as s
        ,       {'NULL': 1}     as f
        ,       1               as i
        ;
    """)
    assert res['ok'] is True, res
    assert {c['name']: c['type'] for c in res['columns']} == {
        'a': 'VARCHAR',
        'b': 'VARCHAR[]',
        's': 'STRUCT(k VARCHAR)',
        'f': 'STRUCT("NULL" INTEGER)',   # a field named NULL is left alone
        'i': 'INTEGER',
    }
    assert _types(res['result_table'])['a'] == 'VARCHAR'


def test_sql_run_ddl_retypes_users_own_table():
    assert sql.run('create table d2_user_null as select null as x, 1 as id;')['ok']
    assert _types('d2_user_null') == {'x': 'VARCHAR', 'id': 'INTEGER'}
    res = sql.run("insert into d2_user_null values ('hello', 2);")
    assert res['ok'] is True, res
    (n,) = db.conn().execute("select count(x) from d2_user_null;").fetchone()
    assert n == 1


def test_sweep_covers_this_cursors_temp_tables():
    c = db.conn()
    c.execute("""
        create temp table d2_temp_null as select null as x;
        create temp view d2_temp_null_view as select null as x;
    """)
    db.retype_null_columns(c)
    types = dict(c.execute("""
        select  table_name
        ,       data_type
        from    duckdb_columns()
        where   database_name = 'temp'
        and     table_name in ('d2_temp_null', 'd2_temp_null_view')
        ;
    """).fetchall())
    assert types == {'d2_temp_null': 'VARCHAR', 'd2_temp_null_view': '"NULL"'}
    c.execute("insert into d2_temp_null values ('text');")


def test_sql_run_leaves_views_alone():
    res = sql.run('create view d2_null_view as select null as x;')
    assert res['ok'] is True, res
    assert _types('d2_null_view') == {'x': '"NULL"'}


def test_save_as_table_retypes():
    res = sql.save_as_table('select null as x', 'd2_saved_null')
    assert res == {'ok': True, 'name': 'd2_saved_null', 'row_count': 1}
    assert _types('d2_saved_null') == {'x': 'VARCHAR'}
    db.conn().execute("insert into d2_saved_null values ('text');")


def test_export_source_query_retypes():
    name, cleanup = ExportSource(kind='query', sql='select null as x').materialize()
    try:
        assert _types(name) == {'x': 'VARCHAR'}
    finally:
        cleanup()


def test_materialize_table_retypes():
    db.conn().execute('create view d2_mat_null as select null as x;')
    assert Api().materialize_table('d2_mat_null') == {'ok': True}
    assert _types('d2_mat_null') == {'x': 'VARCHAR'}
    db.conn().execute("insert into d2_mat_null values ('text');")


def test_ddl_for_null_column_matches_all_null_varchar():
    sql.save_as_table('select null as x', 'd2_ddl_null')
    db.conn().execute('create view d2_ddl_null_view as select null as x;')
    for name in ('d2_ddl_null', 'd2_ddl_null_view'):
        assert _ddl_columns(name, 'sqlserver') == [['x', 'varchar(1)', 'null']]
        assert _ddl_columns(name, 'duckdb') == [['x', 'varchar', 'null']]


@pytest.mark.parametrize('kind', ['flat', 'parquet', 'excel'])
def test_export_query_with_null_column(tmp_path, kind):
    path = tmp_path / f'o.{_EXT[kind]}'
    source = {'kind': 'query', 'name': None, 'sql': 'select 1 as id, null as x'}
    res = export_to_file(source, _target(kind, path))
    assert res['ok'] is True, res
    assert res['rows_written'] == 1
    if kind == 'flat':
        assert path.read_text(encoding='utf-8').splitlines() == ['id,x', '1,']
    elif kind == 'parquet':
        rows = db.conn().execute(f"""
            select  *
            from    read_parquet('{path.as_posix()}')
            ;
        """).fetchall()
        assert rows == [(1, None)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = list(wb['Data'].iter_rows(values_only=True))
        wb.close()
        # read-only openpyxl drops the trailing empty cell
        assert rows == [('id', 'x'), (1,)]


# --------------------------------------------------------------------------- #
# TIMESTAMPTZ (needs pytz to reach Python)                                    #
# --------------------------------------------------------------------------- #

@pytest.fixture(scope='module')
def tz_table(tmp_path_factory):
    pq = tmp_path_factory.mktemp('tz') / 'tz.parquet'
    db.conn().execute(f"""
        copy (
            select  cast(i as integer)                                      as id
            ,       cast('2024-01-02 03:04:05+00' as timestamptz) + to_hours(i) as ts
            from    range(3) as t(i)
        ) to '{pq.as_posix()}' (format parquet)
        ;
    """)
    r = ingest.load_file(str(pq))
    assert _types(r.name)['ts'] == 'TIMESTAMP WITH TIME ZONE'
    return r.name


def test_timestamptz_grid_editor_profile(tz_table):
    rows = Api().get_rows(tz_table)['rows']
    assert len(rows) == 3 and all(isinstance(r[-1], str) for r in rows)

    res = sql.run(f'select ts from {db.quote_ident(tz_table)};')
    assert res['ok'] is True, res
    assert len(res['rows']) == 3

    prof = profile.profile_table(tz_table, 'ts')
    assert prof['common']['non_null_count'] == 3


@pytest.mark.parametrize('kind', ['flat', 'parquet', 'excel'])
def test_timestamptz_export(tmp_path, tz_table, kind):
    path = tmp_path / f'o.{_EXT[kind]}'
    res = export_to_file({'kind': 'table', 'name': tz_table, 'sql': None}, _target(kind, path))
    assert res['ok'] is True, res
    assert res['rows_written'] == 3


def test_excel_time_zone_values_match_csv(tmp_path, tz_table):
    """Excel gets the same wall-clock time the CSV export writes, minus the offset."""
    source = {'kind': 'query', 'name': None, 'sql': f"""
        select  ts
        ,       cast('03:04:05+02' as timetz) as tt
        from    {db.quote_ident(tz_table)}
        order by id
    """}
    csv_path, xlsx_path = tmp_path / 'o.csv', tmp_path / 'o.xlsx'
    assert export_to_file(source, _target('flat', csv_path))['ok']
    assert export_to_file(source, _target('excel', xlsx_path))['ok']

    csv_rows = [line.split(',') for line in csv_path.read_text(encoding='utf-8').splitlines()[1:]]
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    xl_rows = list(wb['Data'].iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(xl_rows) == len(csv_rows) == 3
    for (csv_ts, csv_tt), (xl_ts, xl_tt) in zip(csv_rows, xl_rows):
        assert csv_ts.startswith(xl_ts.isoformat(sep=' '))   # '2024-01-01 22:04:05-05'
        assert csv_tt.startswith(xl_tt.isoformat())          # '03:04:05+02'


def test_excel_failure_releases_temp_file(tmp_path):
    from openpyxl.worksheet import _writer
    before = set(_writer.ALL_TEMP_FILES)
    # A control character is one value openpyxl refuses to write.
    source = {'kind': 'query', 'name': None, 'sql': 'select chr(1) as x'}
    res = export_to_file(source, _target('excel', tmp_path / 'o.xlsx'))
    assert res['ok'] is False
    for path in set(_writer.ALL_TEMP_FILES) - before:
        os.remove(path)   # fails on Windows while the sheet still holds it open


# --------------------------------------------------------------------------- #
# Keyword column names in generated SQL                                       #
# --------------------------------------------------------------------------- #

def test_display_ident_quotes_keywords():
    for name in ('when', 'select', 'ORDER', 'group', 'left', 'true', 'null'):
        assert db.display_ident(name) == f'"{name}"'
    # column_name keywords and plain names read fine bare
    for name in ('time', 'values', 'position', 'amount', '_x1'):
        assert db.display_ident(name) == name


def test_derive_with_keyword_columns():
    db.conn().execute("""
        create or replace table d2_kw as
        select  cast(i as varchar)                         as "when"
        ,       cast(date '2024-01-01' + cast(i as integer) as varchar) as "select"
        ,       'name ' || i                               as "order"
        ,       case when i % 2 = 0 then 'true' else 'false' end as "group"
        ,       'lit'                                      as "true"
        from    range(5) as t(i)
        ;
    """)
    specs = [derive.CastSpec(source=n, target=n, cast_to=t) for n, t in (
        ('when', 'BIGINT'), ('select', 'DATE'), ('order', None), ('group', 'BOOLEAN'),
        ('true', None))]

    preview = derive.preview_derive('d2_kw', specs, None)
    assert [c['name'] for c in preview['columns']] == ['when', 'select', 'order', 'group', 'true']
    assert preview['rows'][1] == [1, '2024-01-02', 'name 1', False, 'lit']

    shown = derive.derive_sql('d2_kw', 'd2_kw_typed', specs, None)
    db.conn().execute(shown)
    assert _types('d2_kw_typed') == {'when': 'BIGINT', 'select': 'DATE', 'order': 'VARCHAR',
                                      'group': 'BOOLEAN', 'true': 'VARCHAR'}

    res = derive.auto_derive('d2_kw', 'order')
    assert res['ok'] is True and res['row_count'] == 5
    assert _types('order')['when'] == 'BIGINT'
    (v,) = db.conn().execute('select "true" from "order" limit 1;').fetchone()
    assert v == 'lit'


# --------------------------------------------------------------------------- #
# Esc-to-cancel                                                               #
# --------------------------------------------------------------------------- #

def test_cancel_interrupts_long_query():
    out: dict = {}
    # Runs for hours unless interrupted.
    worker = threading.Thread(
        target=lambda: out.update(sql.run('select sum(i) from range(10000000000000) as t(i);')),
        daemon=True,
    )
    t0 = time.perf_counter()
    worker.start()
    # Keep cancelling until the run returns: a cancel that lands before
    # the query starts has nothing to interrupt.
    while worker.is_alive() and time.perf_counter() - t0 < 3:
        sql.cancel()
        worker.join(0.05)
    assert not worker.is_alive(), 'query was not interrupted within 3 s'
    assert out['ok'] is False
    assert 'interrupt' in out['error'].lower()
    assert sql.run('select 1 as x;')['rows'] == [[1]]
