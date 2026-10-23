"""CSV load ladder: each DuckDB error class leads to the matching retry."""

from __future__ import annotations

from pathlib import Path

import pytest

from file_quacker import db, ingest, sniff

# 30k rows push the bad row past DuckDB's 20,480-row sniff sample, so the
# error comes from the full read instead of the sniffer.
BODY = b'1,x,y\n' * 30_000
# 'name' and 'note' don't share the 1-character width of the values
# below them, so layout detection sees a header.
HEADER = b'id,name,note\n'


def _write(tmp_path: Path, name: str, data: bytes) -> str:
    p = tmp_path / name
    p.write_bytes(data)
    return str(p)


def _fetch(sql: str) -> list[tuple]:
    return db.conn().execute(sql).fetchall()


# --------------------------------------------------------------------------- #
# Error classifier (strings captured from DuckDB 2.0.0)                       #
# --------------------------------------------------------------------------- #

CAPTURED = {
    'invalid_utf8': (
        'Invalid Input Error: CSV Error on Line: 30002\n'
        'Original Line: 2,caf??z\n'
        'Invalid unicode (byte sequence mismatch) detected. This file is not utf-8 encoded.\n'
        '\n'
        "Possible Solution: Set the correct encoding, if available, to read this CSV File (e.g., encoding='UTF-16')\n"
        'Possible Solution: Enable ignore errors (ignore_errors=true) to skip this row\n'
        '\n'
        '  file = data.csv\n'
        '  delimiter = , (Set By User)\n'
    ),
    'not_cp1252': (
        'Invalid Input Error: CSV Error on Line: 2\n'
        'Original Line: a,b,c\n'
        'Invalid unicode (byte sequence mismatch) detected. This file is not cp1252 encoded.\n'
        '\n'
        "Possible Solution: Set the correct encoding, if available, to read this CSV File (e.g., encoding='UTF-16')\n"
    ),
    'not_utf16': (
        'Invalid Input Error: File is not utf-16 encoded\n'
        '\n'
        "LINE 1: create or replace table t as select * from read_csv('data.csv', header=true, all_varchar=true, strict...\n"
        '                                                   ^^^^^^^^^^^^^^^^^^^^^^^^^^^^^^^^\n'
    ),
    'not_latin1': 'Invalid Input Error: File is not latin-1 encoded\n',
    'too_few': (
        'Invalid Input Error: CSV Error on Line: 30002\n'
        'Original Line: 4,5\n'
        'Expected Number of Columns: 3 Found: 2\n'
        'Possible fixes:\n'
        "* Disable the parser's strict mode (strict_mode=false) to allow reading rows that do not comply with the CSV standard.\n"
        '* Enable null padding (null_padding=true) to replace missing values with NULL\n'
        '* Enable ignore errors (ignore_errors=true) to skip this row\n'
        '\n'
        '  file = data.csv\n'
    ),
    'too_many': (
        'Invalid Input Error: CSV Error on Line: 30002\n'
        'Original Line: \n'
        '4,5,6,7\n'
        'Expected Number of Columns: 3 Found: 4\n'
        'Possible fixes:\n'
        '* Enable null padding (null_padding=true) to replace missing values with NULL\n'
        '\n'
        '  file = data.csv\n'
        '  new_line = \\r\\n (Auto-Detected)\n'
    ),
    'null_padding_quoted_newlines': (
        'CSV Error on Line: 2\n'
        ' The parallel scanner does not support null_padding in conjunction with quoted new lines. '
        'Please disable the parallel csv reader with parallel=false\n'
        '  file = data.csv\n'
        '  null_padding = 1\n'
    ),
    'line_size': (
        'Invalid Input Error: CSV Error on Line: 30002\n'
        'Original Line: 1,xxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx\n'
        'Maximum line size of 2000000 bytes exceeded. Actual Size:2100004 bytes.\n'
        '\n'
        'Possible Solution: Change the maximum length size, e.g., max_line_size=2100006\n'
    ),
    'quote': (
        'Invalid Input Error: CSV Error on Line: 30002\n'
        'Original Line: 1,"abc,z\n'
        '2,y,z\n'
        '\n'
        'Value with unterminated quote found.\n'
        '\n'
        'Possible fixes:\n'
        "* Set quote to empty or to a different value (e.g., quote='')\n"
    ),
    'sniff': (
        'Invalid Input Error: Error when sniffing file "data.csv".\n'
        'It was not possible to automatically detect the CSV parsing dialect\n'
        'The search space used was:\n'
        "Delimiter Candidates: ','\n"
        'Possible fixes:\n'
        '* Enable null padding (null_padding=true) to pad missing columns with NULL values\n'
    ),
}


@pytest.mark.parametrize('key, kind', [
    ('invalid_utf8', 'encoding'),
    ('not_cp1252', 'encoding'),
    ('not_utf16', 'encoding'),
    ('not_latin1', 'encoding'),
    ('too_few', 'too_few'),
    ('too_many', 'too_many'),
    ('null_padding_quoted_newlines', 'null_padding_quoted_newlines'),
    ('line_size', 'line_size'),
    ('quote', 'quote'),
    ('sniff', 'sniff'),
])
def test_classifier_on_captured_errors(key, kind):
    assert ingest._classify_csv_error(CAPTURED[key]) == kind


def test_classifier_ignores_error_text_inside_the_echoed_row():
    # A quoted field echoed in "Original Line" can span lines that look
    # like DuckDB's own descriptions; the real one comes last.
    msg = ('Invalid Input Error: CSV Error on Line: 7\n'
           'Original Line: 1,"\n'
           'Value with unterminated quote found.\n'
           'Invalid unicode (byte sequence mismatch) detected. This file is not utf-8 encoded.\n'
           '",3,4\n'
           'Expected Number of Columns: 3 Found: 4\n'
           'Possible fixes:\n'
           '  file = Expected Number of Columns: 9 Found: 1\n')
    assert ingest._classify_csv_error(msg) == 'too_many'


def test_classifier_unknown_error():
    assert ingest._classify_csv_error('IO Error: Cannot open file "x.csv"') == 'other'


# --------------------------------------------------------------------------- #
# Ladder                                                                      #
# --------------------------------------------------------------------------- #

def test_clean_file_loads_on_first_attempt(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'clean.csv', HEADER + BODY))
    assert r.strategy == 'utf-8'
    assert r.row_count == 30_000 and not r.has_rejects


def test_late_short_row_gets_null_padding(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'short.csv', HEADER + BODY + b'4,5\n' + b'1,x,y\n' * 100))
    assert r.strategy == 'utf-8 + null_padding'
    assert r.row_count == 30_101
    assert _fetch(f'select id, name, note from {db.quote_ident(r.name)} where _src_row_num = 30001') \
        == [('4', '5', None)]


def test_short_row_inside_sniff_sample_gets_null_padding(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'short_early.csv', HEADER + b'1,2,3\n4,5\n6,7,8\n'))
    assert r.strategy == 'utf-8 + no-sniff + null_padding'
    assert r.row_count == 3


def test_quoted_newlines_alone_stay_parallel(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'qnl.csv', HEADER + b'1,"x\ny",3\n' + BODY))
    assert r.strategy == 'utf-8'
    assert _fetch(f'select name from {db.quote_ident(r.name)} where _src_row_num = 1') == [('x\ny',)]


def test_quoted_newlines_with_short_rows_go_sequential(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'qnl_short.csv',
                                HEADER + b'1,"x\ny",3\n4,5\n' + BODY))
    assert r.strategy == 'utf-8 + no-sniff + null_padding + sequential'
    rows = _fetch(f'select id, name, note from {db.quote_ident(r.name)} order by _src_row_num limit 2')
    assert rows == [('1', 'x\ny', '3'), ('4', '5', None)]


def test_late_wide_row_goes_to_rejects_without_mojibake(tmp_path):
    body = 'id,name,city\r\n' + ''.join(f'{i},Chloé,São Paulo\r\n' for i in range(1, 30_001))
    body += '30001,Chloé,São Paulo,EXTRA,EXTRA2\r\n30002,Chloé,Zürich\r\n'
    r = ingest.load_file(_write(tmp_path, 'late_wide.csv', body.encode('utf-8')))
    assert r.encoding == 'utf-8'
    assert r.strategy == 'utf-8 + ignore_errors (1 row rejected)'
    assert r.has_rejects and r.row_count == 30_001
    q = db.quote_ident(r.name)
    assert _fetch(f"select count(*) from {q} where name <> 'Chloé'") == [(0,)]
    assert _fetch(f"select city from {q} where id = '30002'") == [('Zürich',)]
    rejects = _fetch(f"""
        select  distinct line, cast(error_type as varchar), csv_line
        from    {db.quote_ident(r.name + '__rejects')}
    """)
    assert len(rejects) == 1
    line, error_type, csv_line = rejects[0]
    assert line == 30_002 and error_type == 'TOO MANY COLUMNS'
    assert '30001,Chloé,São Paulo,EXTRA,EXTRA2' in csv_line


def test_long_line_loads_on_first_attempt(tmp_path):
    # The explicit read buffer lifts DuckDB's 2 MB line limit.
    r = ingest.load_file(_write(tmp_path, 'long.csv',
                                HEADER + BODY + b'1,' + b'x' * 2_100_000 + b',z\n'))
    assert r.strategy == 'utf-8'
    assert _fetch(f'select length(name) from {db.quote_ident(r.name)} where _src_row_num = 30001') \
        == [(2_100_000,)]


def test_late_unterminated_quote_goes_to_rejects(tmp_path):
    # Reading quotes as data would only save the 2 lines the broken
    # record swallowed, too few to give up quoting for.
    r = ingest.load_file(_write(tmp_path, 'unterminated.csv',
                                HEADER + BODY + b'1,"abc,z\n2,y,z\n'))
    assert r.strategy == 'utf-8 + ignore_errors (1 row rejected, spanning 2 lines)'
    assert r.row_count == 30_000
    (csv_line,) = _fetch(f"select distinct csv_line from {db.quote_ident(r.name + '__rejects')}")[0]
    assert csv_line.startswith('1,"abc,z')


def test_fully_quoted_file_keeps_quoting_around_a_broken_line(tmp_path):
    lines = ['"id","name","city"'] + [f'"{i}","Name{i}","City{i}"' for i in range(1, 30)]
    lines[10] = '"10","unterminated,"City10"'
    r = ingest.load_file(_write(tmp_path, 'quote_all.csv', '\n'.join(lines).encode()))
    q = db.quote_ident(r.name)
    assert r.row_count == 28 and r.has_rejects
    assert _fetch(f"""select count(*) from {q} where name like '"%' or id like '"%'""") == [(0,)]


def _ranked(monkeypatch, *scores):
    monkeypatch.setattr(sniff, 'score_delimiters',
                        lambda path, encoding, quote='"': [sniff.DelimiterScore(*s) for s in scores])


def test_single_column_load_tries_the_runner_up_delimiter(tmp_path, monkeypatch):
    _ranked(monkeypatch, (',', 2, 1.0), ('|', 3, 0.95))
    r = ingest.load_file(_write(tmp_path, 'pipes.txt', b'id|name|note\n1|2|3\n4|5|6\n'))
    assert r.strategy == "utf-8 + delimiter '|'"
    assert r.col_count == 4 and r.delimiter == '|'


def test_single_column_load_falls_back_to_duckdb_delimiter(tmp_path, monkeypatch):
    # The runner-up covers too few rows to try it first.
    _ranked(monkeypatch, (',', 2, 1.0), ('|', 3, 0.5))
    r = ingest.load_file(_write(tmp_path, 'pipes2.txt', b'id|name|note\n1|2|3\n4|5|6\n'))
    assert r.strategy == 'utf-8 + auto-delim'
    assert r.col_count == 4 and r.delimiter is None


def test_single_column_file_stays_single_column(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'one_col.csv', b'name\nalpha\nbeta\n'))
    assert r.strategy == 'utf-8'
    assert r.col_count == 2


def _guess_utf8(monkeypatch):
    # Stands in for a wrong encoding guess, so DuckDB reports the bad bytes.
    monkeypatch.setattr(sniff, 'detect_encoding',
                        lambda path: sniff.EncodingGuess('utf-8', False, 'every byte reads as UTF-8'))


def test_single_byte_file_reads_as_detected(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'latin.csv', HEADER + BODY + b'2,caf\xe9,z\n'))
    assert r.encoding == 'latin-1' and r.strategy == 'latin-1'
    assert 'latin-1 reads it' in r.encoding_note
    assert _fetch(f'select name from {db.quote_ident(r.name)} where _src_row_num = 30001') \
        == [('café',)]


def test_decoding_error_steps_to_cp1252(tmp_path, monkeypatch):
    _guess_utf8(monkeypatch)
    r = ingest.load_file(_write(tmp_path, 'late_cp1252.csv',
                                HEADER + BODY + b'2,caf\xe9,z\n'))
    assert r.encoding == 'cp1252' and r.strategy == 'cp1252'
    assert r.encoding_note.endswith('so it is read as cp1252')
    assert _fetch(f'select name from {db.quote_ident(r.name)} where _src_row_num = 30001') \
        == [('café',)]


def test_bytes_cp1252_lacks_read_as_iso_8859_1(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'c1.csv', HEADER + BODY + b'2,x\x81\xe9,z\n'))
    assert r.encoding == 'iso-8859-1' and r.strategy == 'iso-8859-1'
    assert _fetch(f'select name from {db.quote_ident(r.name)} where _src_row_num = 30001') \
        == [('x\x81é',)]


def test_mostly_utf8_file_rejects_only_rows_with_bad_bytes(tmp_path):
    body = 'id,name,city\n' + ''.join(f'{i},Chloé,Zürich\n' for i in range(1, 2001))
    data = body.encode() + b'2001,caf\xe9,x\n2002,Chlo\xc3\xa9,Z\xfcrich\n' + '2003,Ana,Köln\n'.encode()
    r = ingest.load_file(_write(tmp_path, 'mostly_utf8.csv', data))
    assert r.encoding == 'utf-8'
    assert r.strategy == 'utf-8 + ignore_errors (2 rows rejected)'
    assert 'bad byte sequence' in r.encoding_note
    q = db.quote_ident(r.name)
    assert r.row_count == 2001
    assert _fetch(f"select count(*) from {q} where name = 'Chloé'") == [(2000,)]
    assert _fetch(f"select city from {q} where id = '2003'") == [('Köln',)]
    assert _fetch(f"""
        select  distinct cast(error_type as varchar)
        from    {db.quote_ident(r.name + '__rejects')}
    """) == [('INVALID ENCODING',)]


def test_bad_byte_in_header_keeps_the_data_utf8(tmp_path):
    # DuckDB would drop an undecodable header line and promote row 1, so
    # FQ names the columns itself and DuckDB reads the rows as UTF-8.
    body = ''.join(f'{i},Chloé,Zürich\n' for i in range(1, 2001)).encode()
    r = ingest.load_file(_write(tmp_path, 'bad_header.csv',
                                b'id,n\xe9m,city\n' + body + b'2001,caf\xe9,x\n'))
    assert r.encoding == 'utf-8'
    assert r.strategy == 'utf-8 + own header + ignore_errors (1 row rejected)'
    assert [c['name'] for c in r.columns] == ['_src_row_num', 'id', 'ném', 'city']
    assert r.row_count == 2000
    q = db.quote_ident(r.name)
    assert _fetch(f"""select count(*) from {q} where "ném" = 'Chloé' and city = 'Zürich'""") == [(2000,)]
    assert _fetch(f'select min(_src_row_num), max(_src_row_num), max(cast(id as int)) from {q}') \
        == [(1, 2000, 2000)]
    assert 'header line' in r.encoding_note


def test_bad_byte_in_header_with_short_rows(tmp_path):
    body = ''.join(f'{i},Chloé,Zürich\n' for i in range(1, 2001)).encode()
    r = ingest.load_file(_write(tmp_path, 'bad_header_short.csv',
                                b'id,n\xe9m,city\n' + body + b'2001,Ana\n' + body + b'9,caf\xe9,x\n'))
    assert r.strategy == 'utf-8 + own header + null_padding + ignore_errors (1 row rejected)'
    assert _fetch(f'select "ném", city from {db.quote_ident(r.name)} where _src_row_num = 2001') \
        == [('Ana', None)]


def test_big_endian_utf16_is_refused(tmp_path):
    p = _write(tmp_path, 'be.csv', b'\xfe\xff' + 'id,name\n1,x\n'.encode('utf-16-be'))
    with pytest.raises(ingest.IngestError, match='big-endian UTF-16'):
        ingest.load_file(p)


def test_hash_at_line_start_is_data_not_a_comment(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'hash.csv', b'code,n\n#N/A,1\n#2,2\nx,3\n'))
    assert _fetch(f'select code from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('#N/A',), ('#2',), ('x',)]


def test_user_set_encoding_is_never_changed(tmp_path):
    p = _write(tmp_path, 'forced.csv', HEADER + BODY + b'2,caf\xe9,z\n')
    r = ingest.load_file(p, ingest.IngestOptions(encoding='utf-8'))
    assert r.encoding == 'utf-8'
    assert r.strategy == 'utf-8 (user-set) + ignore_errors (1 row rejected)'
    assert _fetch(f"""
        select  cast(error_type as varchar)
        from    {db.quote_ident(r.name + '__rejects')}
    """) == [('INVALID ENCODING',)]


def test_rejected_undecodable_rows_switch_the_encoding(tmp_path, monkeypatch):
    # ignore_errors (for the wide row) would also hide the cp1252 byte.
    _guess_utf8(monkeypatch)
    data = (HEADER + BODY + b'4,5,6,7\n' + b'1,x,y\n' * 10_000 + b'2,caf\xe9,z\n')
    r = ingest.load_file(_write(tmp_path, 'wide_and_cp1252.csv', data))
    assert r.encoding == 'cp1252'
    assert r.strategy == 'cp1252 + ignore_errors (1 row rejected)'
    assert r.row_count == 40_001


def test_short_rows_rejected_by_user_ignore_errors_get_padded(tmp_path):
    p = _write(tmp_path, 'short_ignore.csv', HEADER + BODY + b'4,5\n' + b'1,x,y\n' * 100)
    r = ingest.load_file(p, ingest.IngestOptions(ignore_errors=True))
    assert r.strategy == 'utf-8 + null_padding + ignore_errors (0 rows rejected)'
    assert r.row_count == 30_101 and not r.has_rejects


def test_every_attempt_listed_when_all_fail(tmp_path):
    p = _write(tmp_path, 'odd_utf16.csv', 'a,b\n1,x\n'.encode('utf-16') + b'\x00')
    with pytest.raises(ingest.IngestError) as info:
        ingest.load_file(p, ingest.IngestOptions(encoding='utf-16', preferred_table_name='odd_utf16'))
    msg = str(info.value)
    assert msg.startswith('Could not read this file after trying ')
    assert '[utf-16 (user-set)] Invalid Input Error: File is not utf-16 encoded' in msg
    assert '[utf-16 (user-set) + null_padding + no-quote + strict=false + ignore_errors]' in msg
    assert _fetch("""
        select  count(*)
        from    information_schema.tables
        where   table_name = 'odd_utf16'
    """) == [(0,)]


def test_view_load_parses_up_front(tmp_path):
    # Views parse lazily, so the ladder reads them once to see errors.
    p = _write(tmp_path, 'view_short.csv', HEADER + BODY + b'4,5\n')
    r = ingest.load_file(p, ingest.IngestOptions(materialize=False))
    assert not r.materialized
    assert r.strategy == 'utf-8 + null_padding'
    assert _fetch(f'select max(_src_row_num) from {db.quote_ident(r.name)}') == [(30_001,)]


def test_caller_reader_options_are_kept(tmp_path):
    p = _write(tmp_path, 'opts.csv', HEADER + BODY)
    r = ingest.load_file(p, ingest.IngestOptions(parallel=False, strict_mode=False, sample_size=-1))
    assert r.strategy == 'utf-8 + sequential + strict=false'
    assert r.row_count == 30_000


# --------------------------------------------------------------------------- #
# Layout: title lines, header row, summary rows                               #
# --------------------------------------------------------------------------- #

REPORT = (b'Monthly report\nGenerated 2024-05-01\n' + HEADER
          + b''.join(b'%d,n%d,note %d\n' % (i, i, i) for i in range(1, 501))
          + b'Total rows: 500\n')


def test_title_lines_and_summary_row_are_detected(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'report.csv', REPORT))
    assert (r.first_row, r.footer_rows) == (3, 1)
    assert [c['name'] for c in r.columns] == ['_src_row_num', 'id', 'name', 'note']
    assert r.row_count == 500
    assert '2 title line(s)' in r.layout_note and 'summary row' in r.layout_note
    q = db.quote_ident(r.name)
    assert _fetch(f'select min(_src_row_num), max(_src_row_num), max(cast(id as int)) from {q}') \
        == [(1, 500, 500)]


def test_plain_file_has_no_layout_note(tmp_path):
    r = ingest.load_file(_write(tmp_path, 'plain.csv', HEADER + BODY))
    assert (r.first_row, r.footer_rows, r.layout_note, r.encoding_note) == (1, 0, None, None)


def test_explicit_first_row_and_skip_footer_are_honored(tmp_path):
    p = _write(tmp_path, 'report2.csv', REPORT)
    r = ingest.load_file(p, ingest.IngestOptions(first_row=3, skip_footer=0))
    assert (r.first_row, r.footer_rows, r.layout_note) == (3, 0, None)
    assert r.row_count == 501
    assert _fetch(f'select id, name from {db.quote_ident(r.name)} where _src_row_num = 501') \
        == [('Total rows: 500', None)]
    r = ingest.load_file(p, ingest.IngestOptions(first_row=0, skip_footer=2))
    assert (r.first_row, r.footer_rows) == (0, 2)
    assert [c['name'] for c in r.columns][:2] == ['_src_row_num', 'col_0']
    assert _fetch(f'select col_0 from {db.quote_ident(r.name)} where _src_row_num = 1') \
        == [('Monthly report',)]
    assert r.row_count == 504 - 2


def test_title_lines_without_header_stay_in_as_rows(tmp_path):
    # DuckDB's sniffer gives up on the narrow title lines, so the ladder
    # hands it the dialect (no-sniff) and pads the short rows.
    data = b'Monthly report\nGenerated 2024-05-01\n' + b''.join(
        b'%d,n%d,%d.5\n' % (i, i, i) for i in range(1, 501))
    r = ingest.load_file(_write(tmp_path, 'titles_no_header.csv', data))
    assert r.first_row == 0 and 'stay in' in r.layout_note
    assert r.strategy == 'utf-8 + no-sniff + null_padding'
    assert r.row_count == 502
    assert _fetch(f'select col_0, col_1 from {db.quote_ident(r.name)} order by _src_row_num limit 3') \
        == [('Monthly report', None), ('Generated 2024-05-01', None), ('1', 'n1')]


def test_preview_reports_detection(tmp_path):
    p = _write(tmp_path, 'report3.csv', REPORT)
    out = ingest.preview_file_lines(str(p), n=5)
    assert out['lines'][:3] == ['Monthly report', 'Generated 2024-05-01', 'id,name,note']
    assert out['encoding'] == 'utf-8' and out['encoding_note'] is None
    assert out['sniffed_delimiter'] == ','
    assert (out['detected_first_row'], out['detected_footer_rows']) == (3, 1)
    assert 'title line' in out['layout_note']


def test_preview_with_a_chosen_encoding(tmp_path):
    p = _write(tmp_path, 'latin2.csv', b'id,name\n1,caf\xe9\n')
    assert ingest.preview_file_lines(str(p))['encoding_note'].startswith('byte 0xE9')
    out = ingest.preview_file_lines(str(p), encoding='cp1252')
    assert out['encoding'] == 'cp1252' and out['encoding_note'] is None
    assert out['lines'] == ['id,name', '1,café']


# --------------------------------------------------------------------------- #
# Row numbers                                                                 #
# --------------------------------------------------------------------------- #

def test_row_numbers_follow_file_order_across_row_groups(tmp_path):
    # 300k rows span three 122,880-row groups and load in parallel.
    data = b'id,v\n' + b''.join(b'%d,v%d\n' % (i, i % 7) for i in range(1, 300_001))
    r = ingest.load_file(_write(tmp_path, 'many_rows.csv', data))
    q = db.quote_ident(r.name)
    assert r.columns[0]['name'] == '_src_row_num'
    assert _fetch(f'select count(*) from {q} where _src_row_num <> cast(id as bigint)') == [(0,)]
    assert _fetch(f'select min(_src_row_num), max(_src_row_num) from {q}') == [(1, 300_000)]


def test_row_numbers_skip_nothing_around_rejected_rows(tmp_path):
    data = (b'id,v\n' + b''.join(b'%d,a\n' % i for i in range(1, 150_001))
            + b'0,a,EXTRA\n' + b''.join(b'%d,a\n' % i for i in range(150_001, 300_001)))
    r = ingest.load_file(_write(tmp_path, 'rownum_rejects.csv', data))
    assert r.strategy == 'utf-8 + ignore_errors (1 row rejected)'
    assert _fetch(f"""
        select  count(*)
        from    {db.quote_ident(r.name)}
        where   _src_row_num <> cast(id as bigint)
    """) == [(0,)]


def test_parquet_row_numbers_follow_file_order(tmp_path):
    p = tmp_path / 'rows.parquet'
    db.conn().execute(f"""
        copy (select range + 1 as id, cast(range % 3 as varchar) as v from range(300000))
        to '{p.as_posix()}' (format parquet, row_group_size 50000)
        ;
    """)
    r = ingest.load_file(str(p))
    q = db.quote_ident(r.name)
    assert r.columns[0]['name'] == '_src_row_num'
    assert _fetch(f'select count(*) from {q} where _src_row_num <> id') == [(0,)]


# --------------------------------------------------------------------------- #
# Placeholder tokens (NA, NaN)                                                #
# --------------------------------------------------------------------------- #

def test_na_nulled_only_where_the_rest_is_numbers_or_dates(tmp_path):
    data = ('amount,country,when,mixed,all_na,lower\n'
            '1.5,US,2024-01-02,1,NA,1\n'
            'NA,NA,NA,NA,NA,na\n'
            'NaN,DE,01/02/2024 3:04 PM,abc,NA,2\n'
            '-2e3,NA,2024-01-03T10:00:00,2,NA,3\n')
    r = ingest.load_file(_write(tmp_path, 'na.csv', data.encode()))
    rows = _fetch(f"""
        select  amount, country, "when", mixed, all_na, lower
        from    {db.quote_ident(r.name)}
        order   by _src_row_num
    """)
    assert [row[0] for row in rows] == ['1.5', None, None, '-2e3']
    assert [row[1] for row in rows] == ['US', 'NA', 'DE', 'NA']
    assert [row[2] for row in rows] == ['2024-01-02', None, '01/02/2024 3:04 PM', '2024-01-03T10:00:00']
    assert [row[3] for row in rows] == ['1', 'NA', 'abc', '2']
    assert [row[4] for row in rows] == ['NA'] * 4
    assert [row[5] for row in rows] == ['1', 'na', '2', '3']


def test_custom_tokens_and_disabling(tmp_path):
    # '#N/A' is kept off the line start: DuckDB's sniffer may take '#' as
    # a comment character.
    p = _write(tmp_path, 'tokens.csv', b'w,v\nNA,5\nx,#N/A\ny,7\n')
    r = ingest.load_file(p, ingest.IngestOptions(nullify_tokens=['#N/A']))
    assert _fetch(f'select v from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('5',), (None,), ('7',)]
    r = ingest.load_file(p, ingest.IngestOptions(nullify_tokens=[]))
    assert _fetch(f'select v from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('5',), ('#N/A',), ('7',)]


def test_parquet_string_columns_use_the_same_rule(tmp_path):
    p = tmp_path / 'na.parquet'
    db.conn().execute(f"""
        copy (
            select  *
            from    (values ('1', 'US', 1), ('NA', 'NA', 2), ('3', 'DE', 3)) v(num, country, n)
        ) to '{p.as_posix()}' (format parquet)
        ;
    """)
    r = ingest.load_file(str(p))
    assert _fetch(f'select num, country from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('1', 'US'), (None, 'NA'), ('3', 'DE')]


def test_xlsx_uses_the_same_rule(tmp_path):
    import openpyxl
    p = tmp_path / 'na.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['num', 'country'])
    ws.append([1.25, 'US'])
    ws.append(['NA', 'NA'])
    ws.append([11 / 3, 'DE'])
    wb.save(p)
    r = ingest.load_file(str(p))
    assert _fetch(f'select num, country from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('1.25', 'US'), (None, 'NA'), ('3.66666667', 'DE')]


@pytest.mark.parametrize('stats, tokens, expected', [
    # numbers only, token above the max: no token can be there
    ({'has_no_null': True, 'min': '-1', 'max': '99'}, [b'NA'], False),
    # numbers plus the token as max
    ({'has_no_null': True, 'min': '-1', 'max': 'NaN'}, [b'NA', b'NaN'], True),
    # text column
    ({'has_no_null': True, 'min': 'Ana', 'max': 'Zed'}, [b'NA'], False),
    # max is a cut-off prefix of a long token
    ({'has_no_null': True, 'min': '1', 'max': 'NaNaNaNaNaNa'}, [b'NaNaNaNaNaNaNaN'], True),
    # token sorting below the digits
    ({'has_no_null': True, 'min': '#N/A', 'max': '9'}, [b'#N/A'], True),
    # escaped backslash in stats
    ({'has_no_null': True, 'min': '-1', 'max': '\\x5CN'}, [b'\\N'], True),
    # all null
    ({'has_no_null': False, 'has_null': True}, [b'NA'], False),
])
def test_stats_precheck(stats, tokens, expected):
    assert ingest._may_hold_only_tokens_and_shapes(stats, tokens) is expected


def test_xlsx_date_and_time_columns_convert_in_one_pass(tmp_path):
    import datetime as dt
    import openpyxl
    p = tmp_path / 'dates.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['day', 'at', 'n'])
    for day, at, n in [(dt.date(2024, 1, 2), dt.time(12, 0), 0.1 + 0.2),
                       (dt.date(2024, 2, 29), dt.time(6, 30, 15), 5)]:
        ws.append([day, at, n])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = 'yyyy-mm-dd'
        row[1].number_format = 'hh:mm:ss'
    wb.save(p)
    r = ingest.load_file(str(p))
    assert _fetch(f'select day, "at", n from {db.quote_ident(r.name)} order by _src_row_num') == [
        ('2024-01-02', '12:00:00', '0.30'),
        ('2024-02-29', '06:30:15', '5'),
    ]


# --------------------------------------------------------------------------- #
# End-to-end punch list                                                       #
# --------------------------------------------------------------------------- #

def test_utf16le_spelling_reads_as_duckdb_utf16(tmp_path):
    p = _write(tmp_path, 'le.csv', 'id,name\n1,Zoë\n2,Ångström\n'.encode('utf-16-le'))
    r = ingest.load_file(p, ingest.IngestOptions(encoding='utf-16le'))
    assert r.strategy == 'utf-16 (user-set)'
    assert _fetch(f'select id, name from {db.quote_ident(r.name)} order by _src_row_num') \
        == [('1', 'Zoë'), ('2', 'Ångström')]


def test_unsupported_encoding_stops_with_a_short_message(tmp_path):
    p = _write(tmp_path, 'be2.csv', 'id,name\n1,x\n'.encode('utf-16-be'))
    with pytest.raises(ingest.IngestError) as info:
        ingest.load_file(p, ingest.IngestOptions(encoding='utf-16be'))
    assert str(info.value) == ("DuckDB cannot read the encoding 'utf-16be'. "
                               'Pick another one, or leave the encoding on Auto.')


def test_error_lines_keep_the_problem_and_drop_the_noise(tmp_path):
    p = _write(tmp_path, 'odd2.csv', 'a,b\n1,x\n'.encode('utf-16') + b'\x00')
    with pytest.raises(ingest.IngestError) as info:
        ingest.load_file(p, ingest.IngestOptions(encoding='utf-16'))
    attempts = str(info.value).splitlines()[1:]
    assert attempts and all(len(a) < 450 for a in attempts)
    assert all('File is not utf-16 encoded' in a for a in attempts)
    assert 'LINE ' not in str(info.value)
    detail = ingest._error_detail(Exception(CAPTURED['too_many']))
    assert detail == ('Invalid Input Error: CSV Error on Line: 30002 | Original Line: | 4,5,6,7 | '
                      'Expected Number of Columns: 3 Found: 4')
    assert 'search space' not in ingest._error_detail(Exception(CAPTURED['sniff']))


def test_preview_of_utf16_without_bom(tmp_path):
    p = _write(tmp_path, 'nobom.csv', ('id,name\n' + ''.join(f'{i},n{i}\n' for i in range(50))).encode('utf-16-le'))
    out = ingest.preview_file_lines(str(p), n=3)
    assert out['encoding'] == 'utf-16' and out['lines'] == ['id,name', '0,n0', '1,n1']


def test_unterminated_quote_among_quoted_fields(tmp_path):
    rows = [f'{i},"Smith, J{i}",c' for i in range(1, 30)]
    rows[9] = '10,"unterminated,c'
    r = ingest.load_file(_write(tmp_path, 'quoted_bad.csv', ('id,name,city\n' + '\n'.join(rows) + '\n').encode()))
    # Rows 10 and 11 read as one broken record; the other quoted commas stay.
    assert r.strategy == 'utf-8 + no-sniff + ignore_errors (1 row rejected, spanning 2 lines)'
    q = db.quote_ident(r.name)
    assert r.row_count == 27
    assert _fetch(f"select count(*) from {q} where name like 'Smith, J%' and city = 'c'") == [(27,)]
    (csv_line,) = _fetch(f"select distinct csv_line from {db.quote_ident(r.name + '__rejects')}")[0]
    assert csv_line.startswith('10,"unterminated,c')


def test_quote_that_never_closes_reads_quotes_as_data(tmp_path):
    rows = ''.join(f'{i},n{i},c{i}\n' if i != 10 else '10,"open quote,c10\n' for i in range(1, 30))
    r = ingest.load_file(_write(tmp_path, 'runaway.csv', ('id,name,city\n' + rows).encode()))
    assert r.strategy == 'utf-8 + no-sniff + no-quote + ignore_errors (0 rows rejected)'
    # The footer was detected with quotes on, so it isn't applied here.
    assert (r.row_count, r.footer_rows) == (29, 0)
    assert _fetch(f'select name from {db.quote_ident(r.name)} where id = \'10\'') == [('"open quote',)]


def test_no_sniff_names_match_duckdb(tmp_path):
    header = 'id,, a ,a,A,"q,x",column1\n'
    clean = ingest.load_file(_write(tmp_path, 'names_clean.csv', (header + '1,2,3,4,5,6,7\n' * 5).encode()))
    broken = ingest.load_file(_write(tmp_path, 'names_broken.csv',
                                     (header + '1,2,3,4,5,6,7\n' * 5 + '1,"2,3,4,5,6,7\n').encode()))
    assert 'no-sniff' in broken.strategy
    assert [c['name'] for c in broken.columns] == [c['name'] for c in clean.columns]


def test_parquet_null_column_becomes_varchar(tmp_path):
    p = tmp_path / 'nulls.parquet'
    db.conn().execute(f"copy (select 1 as id, null as nothing) to '{p.as_posix()}' (format parquet)")
    r = ingest.load_file(str(p))
    assert {c['name']: c['type'] for c in r.columns}['nothing'] == 'VARCHAR'


def test_cleared_number_fields_mean_auto(tmp_path):
    p = _write(tmp_path, 'report4.csv', REPORT)
    r = ingest.load_file(p, ingest.IngestOptions(first_row='', skip_footer='', skip_rows=''))
    assert (r.first_row, r.footer_rows, r.row_count) == (3, 1, 500)


def test_preview_redetects_layout_for_a_chosen_delimiter(tmp_path):
    data = b'id|name|note\n' + b''.join(b'%d|a,b|c\n' % i for i in range(1, 200)) + b'total|9|\n'
    p = _write(tmp_path, 'pipes3.txt', data)
    auto = ingest.preview_file_lines(str(p))
    assert auto['sniffed_delimiter'] == '|'
    forced = ingest.preview_file_lines(str(p), delimiter=',')
    assert forced['sniffed_delimiter'] == ','
    assert forced['layout_note'] != auto['layout_note']


def test_big_file_with_long_first_lines_numbers_from_rowid(tmp_path, monkeypatch):
    # Long lines at the top would talk the row estimate down; a file bigger
    # than the inline buffer always takes the rowid path.
    p = Path(_write(tmp_path, 'long_head.csv', b'id,v\n' + b'1,' + b'x' * 4000 + b'\n' + b'2,y\n' * 50))
    monkeypatch.setattr(ingest, 'INLINE_READ_BUFFER_BYTES', 1024)
    assert not ingest._fits_inline(p)
    monkeypatch.setattr(ingest, 'INLINE_READ_BUFFER_BYTES', 1 << 20)
    assert ingest._fits_inline(p)
